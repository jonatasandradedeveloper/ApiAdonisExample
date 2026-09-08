#!/usr/bin/env python3
"""Gera o documento Word ÚNICO da Andressa Silva (capa + sumário + todas as seções)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from generate_complete_docs import (
    GRAY,
    NAVY,
    add_footer,
    add_runs,
    markdown_to_doc,
    set_cell_shading,
    set_run_font,
    style_doc,
)

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "word"
AND_DIR = ROOT / "andressa"

OUTPUT = OUT / "Andressa-Silva-Documento-Unico-People-Lead.docx"

SECTIONS = [
    ("00-NOTA-uso-interno-pl.md", "Nota — uso interno (não enviar à Andressa)"),
    ("00-ficha.md", "1. Ficha e contexto"),
    ("07-resumo-call-1.md", "2. Resumo da 1:1 #1 (28/07/2026)"),
    ("01-historico-1-1.md", "3. Histórico completo de 1:1s"),
    ("02-prioridades.md", "4. Prioridades e parecer do People Lead"),
    ("03-plano-carreira.md", "5. Plano de carreira"),
    ("06-planejamento-carreira-empresa.md", "6. Plano de carreira × o que a empresa precisa"),
    ("04-insumos-ciclo.md", "7. Insumos para o December Cycle"),
    ("05-acoes.md", "8. Ações e checklist"),
    ("14-pedidos-feedback.md", "9. Pedidos de feedback (gestor / colega)"),
    ("16-feedbacks-projeto.md", "10. Feedbacks de projeto (Camila, Renata, Rafael)"),
    ("15-justificativa-recomendacao.md", "11. Justificativa da Recomendação"),
    ("08-workday-prioridades-abcd.md", "12. Workday — prioridades e ABCD"),
    ("09-autorreflexao-passo-a-passo.md", "13. Autorreflexão ABCD — referência PL"),
    ("13-orientacao-ciclo-andressa.md", "14. Orientação do ciclo (Andressa Silva)"),
]

VEREDITO = (
    "Analyst MAPFRE: WIT (Camila) e par Renata (C#/SQL) fortes. "
    "Gestor Rafael G. Martins: Review in 12 months; evolução técnica abaixo do esperado; "
    "foco autonomia + .NET. Não promover neste FY. Justificativa pronta."
)


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()

    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = eyebrow.add_run("AVANADE · PEOPLE LEAD · DOCUMENTO ÚNICO")
    set_run_font(r, size=12, bold=True, color=NAVY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Andressa Silva")
    set_run_font(r, size=28, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Acompanhamento completo · carreira · prioridades · Workday · December Cycle FY26"
    )
    set_run_font(r, size=12, italic=True, color=GRAY)

    doc.add_paragraph()

    info = doc.add_table(rows=7, cols=2)
    info.style = "Table Grid"
    rows = [
        ("Cargo", "Analyst, Back-End Developer"),
        ("Cliente / projeto", "MAPFRE — sustentação (stack legado)"),
        ("Gestor projeto", "Rafael G. Martins — Review in 12 months (feedback 08/09)"),
        ("People Lead", "Jônatas Andrade Da Silva"),
        ("Na Avanade desde", "2022"),
        ("Prazos", "Justificativa pronta · input PL em setembro"),
        ("Última atualização", "8 de setembro de 2026"),
    ]
    for i, (k, v) in enumerate(rows):
        pk = info.rows[i].cells[0].paragraphs[0]
        pv = info.rows[i].cells[1].paragraphs[0]
        add_runs(pk, k, bold=True)
        add_runs(pv, v)
        set_cell_shading(info.rows[i].cells[0], "F3F6FA")

    doc.add_paragraph()
    box = doc.add_paragraph()
    r = box.add_run("Síntese do People Lead")
    set_run_font(r, size=12, bold=True, color=NAVY)
    p = doc.add_paragraph()
    add_runs(p, VEREDITO)

    note = doc.add_paragraph()
    r = note.add_run(
        "Histórico interno do People Lead sobre Andressa Silva — NÃO enviado à profissional. "
        "Textos de Workday/ABCD são referência sua (orientação na call; cobrança; input PL). "
        "Atualizar após cada 1:1."
    )
    set_run_font(r, size=10, italic=True, color=GRAY)

    doc.add_page_break()


def add_toc(doc: Document) -> None:
    doc.add_heading("Sumário", level=1)
    for _, title in SECTIONS:
        p = doc.add_paragraph()
        add_runs(p, title)
        p.paragraph_format.space_after = Pt(6)

    tip = doc.add_paragraph()
    r = tip.add_run(
        "Dica: no Word, use Referências → Sumário para índice com números de página."
    )
    set_run_font(r, size=9, italic=True, color=GRAY)
    doc.add_page_break()


def build() -> Path:
    doc = Document()
    style_doc(doc)
    add_footer(doc, "Andressa Silva")
    add_cover(doc)
    add_toc(doc)

    for idx, (filename, section_title) in enumerate(SECTIONS):
        path = AND_DIR / filename
        if not path.exists():
            print(f"SKIP  {filename} (not found)")
            continue
        if idx > 0:
            doc.add_page_break()
        md = path.read_text(encoding="utf-8")
        doc.add_heading(section_title, level=1)
        markdown_to_doc(md, doc, skip_first_h1=True, demote=1)

    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


def main() -> None:
    path = build()
    print(f"OK  {path.relative_to(ROOT)}  ({path.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
